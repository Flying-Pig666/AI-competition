# -*- coding: utf-8 -*-
"""生成两份PPT:
1. 驼研信鉴_视频版_自动换片.pptx  —— 每页已写入自动换片时间,合计90秒,按F5播放即可录屏
2. 驼研信鉴_项目介绍版.pptx      —— 16页完整介绍,手动翻页,答辩/路演用
深蓝科技风,流程图用圆角矩形+箭头直接绘制。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 配色(深蓝科技风) ----------
BG    = RGBColor(0x0B, 0x1B, 0x33)   # 主背景 深蓝黑
PANEL = RGBColor(0x14, 0x2A, 0x4A)   # 卡片底
PANEL2= RGBColor(0x1B, 0x35, 0x5C)   # 亮一档卡片
LINE  = RGBColor(0x2E, 0x5F, 0x8A)   # 描边
GOLD  = RGBColor(0xF0, 0xB4, 0x29)   # 金色强调
CYAN  = RGBColor(0x4C, 0xC9, 0xF0)   # 科技蓝
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUB   = RGBColor(0x9F, 0xB3, 0xC8)   # 次要文字
RED   = RGBColor(0xE4, 0x57, 0x2E)   # 痛点红
GREEN = RGBColor(0x2E, 0xC4, 0xB6)   # 成果绿

FONT = "微软雅黑"
W, H = 13.333, 7.5
OUT_DIR = r"C:\AI金融赛\提交材料"
QR_PNG = os.path.join(OUT_DIR, "演示二维码_扫码在线体验.png")
DEMO_URL = "https://ai-competition-sziwvynrh9y4cmqqu3h3fu.streamlit.app"


# ---------- 基础助手 ----------
def _set_run(r, text, size, color, bold=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', FONT)


def tx(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold) 或 (text, size, color, bold, space_after_pt)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, spec in enumerate(lines):
        text, size, color, bold = spec[0], spec[1], spec[2], spec[3]
        space = spec[4] if len(spec) > 4 else 6
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        _set_run(p.add_run(), text, size, color, bold)
    return tb


def bg(slide):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    r.fill.solid(); r.fill.fore_color.rgb = BG
    r.line.fill.background(); r.shadow.inherit = False
    # 顶部装饰金线
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(0.07))
    bar.fill.solid(); bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background(); bar.shadow.inherit = False


def page_title(slide, text, sub=None):
    tx(slide, 0.6, 0.35, 12.1, 1.0, [(text, 32, WHITE, True)])
    if sub:
        tx(slide, 0.6, 1.05, 12.1, 0.5, [(sub, 15, SUB, False)])


def box(slide, x, y, w, h, lines, fill=PANEL, line=LINE, align=PP_ALIGN.CENTER):
    """圆角卡片;lines 同 tx()"""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line; shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, spec in enumerate(lines):
        text, size, color, bold = spec[0], spec[1], spec[2], spec[3]
        space = spec[4] if len(spec) > 4 else 4
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        _set_run(p.add_run(), text, size, color, bold)
    return shp


def arrow_r(slide, x, y, w=0.55, h=0.42, color=GOLD):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def arrow_d(slide, x, y, w=0.42, h=0.5, color=GOLD):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def set_auto_advance(slide, seconds):
    """写入自动换片时间(毫秒),禁止单击换片,淡入淡出过渡"""
    sld = slide._element
    for t in sld.findall(qn('p:transition')):
        sld.remove(t)
    tr = sld.makeelement(qn('p:transition'), {
        'advClick': '0', 'advTm': str(int(seconds * 1000)), 'spd': 'slow'})
    tr.append(tr.makeelement(qn('p:fade'), {}))
    timing = sld.find(qn('p:timing'))
    if timing is not None:
        timing.addprevious(tr)
    else:
        sld.append(tr)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def add_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg(s)
    return s


# ============================================================
# 版本一:视频版(自动换片,合计90秒)
# ============================================================
def build_video():
    prs = new_deck()

    # P1 标题 (7s)
    s = add_slide(prs)
    tx(s, 0, 2.0, W, 1.4, [("驼研 · 信鉴", 66, GOLD, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 3.5, W, 0.9, [("基于公告与财报多源结构化提取的上市公司信用风险预警智能体", 22, WHITE, False)],
       align=PP_ALIGN.CENTER)
    tx(s, 0, 4.45, W, 0.5, [("2026 北京市大学生金融人工智能竞赛", 16, SUB, False)], align=PP_ALIGN.CENTER)
    # 装饰小方块
    for i, c in enumerate([GOLD, CYAN, GREEN]):
        d = slide_dot = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                           Inches(6.17 + i * 0.5 - 0.5), Inches(5.4), Inches(0.28), Inches(0.28))
        d.adjustments[0] = 0.3
        d.fill.solid(); d.fill.fore_color.rgb = c
        d.line.fill.background(); d.shadow.inherit = False
    set_auto_advance(s, 7)

    # P2 三个现实难题 (10s)
    s = add_slide(prs)
    page_title(s, "三个现实难题")
    cards = [
        ("爆雷频发", "大股东高比例质押\n是埋在股价下的地雷", RED),
        ("公告海量", "数千家公司天天发公告\n格式各异,人工读不过来", GOLD),
        ("AI 会编数", "直接问大模型\n它可能一本正经地编数字", CYAN),
    ]
    for i, (t, d, c) in enumerate(cards):
        box(s, 0.7 + i * 4.15, 2.2, 3.75, 3.2,
            [(t, 26, c, True, 14), (d, 16, WHITE, False)])
    set_auto_advance(s, 10)

    # P3 康美案例 (9s)
    s = add_slide(prs)
    tx(s, 0, 1.5, W, 1.6, [("300 亿", 92, RED, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 3.4, W, 0.6, [("康美药业造假爆雷", 28, WHITE, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 4.35, W, 1.2,
       [("爆雷前,大股东的股票早已几乎全部质押", 19, SUB, False, 8),
        ("—— 信号就写在公告里,只是没人读得完", 19, GOLD, False)], align=PP_ALIGN.CENTER)
    set_auto_advance(s, 9)

    # P4 方案总览 (10s)
    s = add_slide(prs)
    page_title(s, "驼研·信鉴:公告进,预警出")
    bx, bw, bh, by = 0.55, 2.15, 1.7, 2.6
    labels = [("质押公告 PDF", "文本层 / 扫描件", CYAN),
              ("结构化提取", "LLM + 校验重试", GOLD),
              ("程序勾稽验算", "错就自动打回", GOLD),
              ("信用风险预警", "红旗规则 R1-R9", GREEN)]
    # 双入口:公告+财报
    box(s, bx, by - 0.95, bw, 0.8, [("质押公告 PDF", 17, CYAN, True)])
    box(s, bx, by + 0.35, bw, 0.8, [("历年财报数据", 17, CYAN, True)])
    for i, (t, d, c) in enumerate(labels[1:], start=1):
        arrow_r(s, bx + bw + 0.08 + (i - 1) * (bw + 0.75), by + 0.42)
        box(s, bx + i * (bw + 0.75), by - 0.3, bw, bh,
            [(t, 19, c, True, 8), (d, 13, SUB, False)])
    tx(s, 0, 5.3, W, 0.5, [("全程留痕 · 每个字段都能找回公告原文出处", 16, SUB, False)],
       align=PP_ALIGN.CENTER)
    set_auto_advance(s, 10)

    # P5 选题一 (11s)
    s = add_slide(prs)
    page_title(s, "选题一 | 公告数据结构化提取")
    steps = [("双通道解析", "文本层PDF\n扫描件OCR"),
             ("LLM 提取", "按标准Schema\n输出JSON"),
             ("校验-重试", "程序验算不过\n自动打回≤3轮"),
             ("质疑-裁决", "两个AI\n回原文互相监督"),
             ("标准 JSON", "12/12校验\n逐字段溯源")]
    bw2, gap = 2.12, 0.35
    x0 = (W - (bw2 * 5 + gap * 4)) / 2
    for i, (t, d) in enumerate(steps):
        xx = x0 + i * (bw2 + gap)
        if i:
            arrow_r(s, xx - gap + 0.02, 2.95, w=0.32, h=0.36, color=CYAN)
        box(s, xx, 2.35, bw2, 1.6,
            [(t, 17, GOLD if i in (2, 3) else WHITE, True, 8), (d, 12.5, SUB, False)])
    tx(s, 0, 4.5, W, 0.9,
       [("提取员干活 · 质疑员挑刺 · 裁决员拍板", 20, WHITE, True, 8),
        ("实测合力发现过“公告自己印错了”的数据矛盾", 15, GREEN, False)], align=PP_ALIGN.CENTER)
    set_auto_advance(s, 11)

    # P6 选题二:九条红旗 (11s)
    s = add_slide(prs)
    page_title(s, "选题二 | 财报分析 · 九条红旗规则")
    rules = [("R1", "利润与现金流背离"), ("R2", "非经常性损益异常"), ("R3", "营收与利润背离"),
             ("R4", "业绩变脸"), ("R5", "持续亏损"), ("R6", "高杠杆"),
             ("R7", "利润含金量不足"), ("R8", "股东高比例质押"), ("R9", "质押到期压力")]
    cw, ch, cg = 3.85, 1.25, 0.28
    gx = (W - (cw * 3 + cg * 2)) / 2
    for i, (rid, name) in enumerate(rules):
        r, c = divmod(i, 3)
        hot = rid in ("R8", "R9")  # 联动选题一的两条
        box(s, gx + c * (cw + cg), 1.55 + r * (ch + cg), cw, ch,
            [(rid + "  " + name, 16, GOLD if hot else WHITE, True)],
            fill=PANEL2 if hot else PANEL, line=GOLD if hot else LINE)
    tx(s, 0, 6.15, W, 0.9,
       [("规则源自 Sloan 1996 · Beneish 1999 · Altman 1968 · 谢德仁等 2016 等经典研究", 15, SUB, False, 6),
        ("金色 R8 / R9 由选题一的公告提取结果直接驱动 —— 两个选题的咬合点", 15, GOLD, False)],
       align=PP_ALIGN.CENTER)
    set_auto_advance(s, 11)

    # P7 三重保障 (9s)
    s = add_slide(prs)
    page_title(s, "为什么敢信它的输出?")
    guards = [("程序算数", "公式全部由代码计算\nAI 只解释,不碰数字", GREEN),
              ("双源印证", "同花顺 × 新浪\n交叉核对,留痕落盘", CYAN),
              ("全程留痕", "每步可追溯\n幻觉字段无处遁形", GOLD)]
    for i, (t, d, c) in enumerate(guards):
        box(s, 0.7 + i * 4.15, 2.2, 3.75, 3.0,
            [(t, 24, c, True, 14), (d, 16, WHITE, False)])
    set_auto_advance(s, 9)

    # P8 实测数字 (9s)
    s = add_slide(prs)
    page_title(s, "30 份未见公告 · 盲测数字")
    nums = [("100%", "跑通率", GREEN), ("93.5%", "校验项通过率", GOLD),
            ("91.5%", "字段溯源覆盖率", CYAN), ("3 倍", "文档级全过率\n纯规则 20% → 本系统 60%", GREEN)]
    cw3 = 2.85
    gx = (W - (cw3 * 4 + 0.35 * 3)) / 2
    for i, (n, d, c) in enumerate(nums):
        box(s, gx + i * (cw3 + 0.35), 2.3, cw3, 2.9,
            [(n, 44, c, True, 12), (d, 15, WHITE, False)])
    set_auto_advance(s, 9)

    # P9 康美回溯时间线 (8s)
    s = add_slide(prs)
    page_title(s, "回到爆雷之前")
    tl = [("2017 年报", "R2 首次报警", GOLD), ("2018 年报", "R1 / R4 连续报警", RED),
          ("2019.4", "造假曝光", SUB)]
    lw = 3.3
    gx = (W - (lw * 3 + 1.1 * 2)) / 2
    for i, (t, d, c) in enumerate(tl):
        if i:
            arrow_r(s, gx + i * (lw + 1.1) - 1.0, 3.0, w=0.9, h=0.4, color=SUB)
        box(s, gx + i * (lw + 1.1), 2.5, lw, 1.5,
            [(t, 20, c, True, 8), (d, 16, WHITE, False)])
    tx(s, 0, 4.7, W, 1.0,
       [("红色预警比爆雷早约 1-2 年", 22, GOLD, True, 8),
        ("贵州茅台等健康公司:零误报", 16, GREEN, False)], align=PP_ALIGN.CENTER)
    set_auto_advance(s, 8)

    # P10 结尾 (6s)
    s = add_slide(prs)
    tx(s, 0, 1.2, W, 1.0, [("公告进,预警出", 44, GOLD, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 2.5, W, 0.5, [("在线演示 · 扫码亲手体验", 20, WHITE, False)], align=PP_ALIGN.CENTER)
    if os.path.exists(QR_PNG):
        s.shapes.add_picture(QR_PNG, Inches(W / 2 - 0.95), Inches(3.1), Inches(1.9), Inches(1.9))
    tx(s, 0, 5.15, W, 0.45, [(DEMO_URL, 14, SUB, False)], align=PP_ALIGN.CENTER)
    set_auto_advance(s, 6)

    path = os.path.join(OUT_DIR, "驼研信鉴_视频版_自动换片.pptx")
    prs.save(path)
    return path, 7 + 10 + 9 + 10 + 11 + 11 + 9 + 9 + 8 + 6


# ============================================================
# 版本二:项目介绍版(16页,手动翻页)
# ============================================================
def build_intro():
    prs = new_deck()

    # P1 封面
    s = add_slide(prs)
    tx(s, 0, 1.6, W, 1.3, [("驼研 · 信鉴", 60, GOLD, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 3.0, W, 0.9, [("基于公告与财报多源结构化提取的上市公司信用风险预警智能体", 22, WHITE, False)],
       align=PP_ALIGN.CENTER)
    tx(s, 0, 4.6, W, 1.4,
       [("2026 北京市大学生金融人工智能竞赛", 16, SUB, False, 8),
        ("参赛学校:________  队名:________  成员:________", 15, SUB, False)],
       align=PP_ALIGN.CENTER)

    # P2 目录
    s = add_slide(prs)
    page_title(s, "目录")
    items = ["01  背景与痛点", "02  总体方案", "03  选题一:公告数据结构化提取",
             "04  选题二:财报分析与信用风险预警", "05  实验验证", "06  创新与总结"]
    for i, it in enumerate(items):
        r, c = divmod(i, 2)
        box(s, 1.2 + c * 5.7, 1.7 + r * 1.75, 5.2, 1.35, [(it, 20, WHITE, True)],
            align=PP_ALIGN.LEFT)

    # P3 背景
    s = add_slide(prs)
    page_title(s, "背景:股权质押是什么,为什么重要")
    box(s, 0.8, 1.7, 5.9, 4.6,
        [("什么是股权质押", 22, GOLD, True, 12),
         ("大股东缺钱了,把手里的股票抵押给金融机构借钱。", 17, WHITE, False, 8),
         ("按规定,每次质押/解押/展期都要发公告,写清谁质押、多少股、何时到期。", 17, WHITE, False)],
        align=PP_ALIGN.LEFT)
    box(s, 7.0, 1.7, 5.5, 4.6,
        [("为什么要盯着它", 22, RED, True, 12),
         ("质押比例过高,往往是公司出事的前兆。", 17, WHITE, False, 8),
         ("康美药业 300 亿造假爆雷前,大股东股票早已几乎全部质押。", 17, WHITE, False, 8),
         ("这些信号对银行、投资者、监管都至关重要。", 17, WHITE, False)],
        align=PP_ALIGN.LEFT)

    # P4 痛点
    s = add_slide(prs)
    page_title(s, "痛点:要把公告用起来,有三道坎")
    pains = [("坎一 · 公告难读", "几千家公司天天发,格式千奇百怪,还有扫描图片;人工提取慢且易错"),
             ("坎二 · 数字不能错", "金融数据错一个数就可能误判,要求接近零容错"),
             ("坎三 · 大模型会编", "AI 直接读公告会编造看似合理的数字,还说不出出处——金融场景致命")]
    for i, (t, d) in enumerate(pains):
        box(s, 0.8, 1.7 + i * 1.7, 11.7, 1.45,
            [(t, 19, RED, True, 6), (d, 16, WHITE, False)], align=PP_ALIGN.LEFT)

    # P5 总体方案
    s = add_slide(prs)
    page_title(s, "总体方案:两个选题,一条数据链")
    box(s, 0.6, 2.0, 2.4, 1.0, [("质押公告", 18, CYAN, True)])
    box(s, 0.6, 3.4, 2.4, 1.0, [("历年财报", 18, CYAN, True)])
    arrow_r(s, 3.15, 2.55, w=0.6, h=0.9, color=CYAN)
    box(s, 3.9, 2.0, 3.1, 2.4,
        [("选题一", 15, SUB, False, 4), ("结构化提取", 20, GOLD, True, 8),
         ("谁质押/多少股/何时到期\n标准 JSON 输出", 14, SUB, False)])
    arrow_r(s, 7.15, 2.9, w=0.6, h=0.6)
    box(s, 7.9, 2.0, 4.8, 2.4,
        [("选题二", 15, SUB, False, 4), ("信用风险预警", 20, GREEN, True, 8),
         ("财报红旗 R1-R7 + 公告驱动 R8/R9\n红色/橙色分级预警报告", 14, SUB, False)])
    tx(s, 0.6, 5.0, 12.1, 1.4,
       [("咬合点:选题一提取的质押比例与到期压力,直接成为选题二的 R8/R9 两条规则", 17, GOLD, True, 8),
        ("三重把关:程序勾稽验算 + 逐字段回原文溯源 + 质疑-裁决多智能体审查", 17, WHITE, False)])

    # P6 选题一·解析
    s = add_slide(prs)
    page_title(s, "选题一(1/3):先读懂公告 —— 双通道解析")
    box(s, 0.8, 1.8, 5.7, 3.9,
        [("通道 A · 文本层 PDF", 20, CYAN, True, 10),
         ("pypdf / pdfplumber 直接抽取文字与表格", 16, WHITE, False, 6),
         ("速度快、零噪声,是首选通道", 15, SUB, False)],
        align=PP_ALIGN.LEFT)
    box(s, 6.8, 1.8, 5.7, 3.9,
        [("通道 B · 扫描件 OCR", 20, GOLD, True, 10),
         ("逐页转图 → RapidOCR 识别", 16, WHITE, False, 6),
         ("用行框坐标重建表格结构", 16, WHITE, False, 6),
         ("OCR 噪声字段会被溯源机制显式标记", 15, SUB, False)],
        align=PP_ALIGN.LEFT)
    tx(s, 0.8, 6.0, 11.7, 0.6,
       [("实测:文本件 12/12 校验通过;扫描件 12/12 通过、溯源覆盖率 86%(噪声字段全部正确标记)", 15, GREEN, False)])

    # P7 选题一·闭环
    s = add_slide(prs)
    page_title(s, "选题一(2/3):AI 提取 + 校验-重试闭环")
    flow = [("LLM 提取", "按 Schema 出 JSON"), ("程序校验", "7 项勾稽验算"),
            ("不过审?", "问题清单回灌"), ("打回重做", "最多 3 轮")]
    bw, gap = 2.6, 0.5
    x0 = (W - (bw * 4 + gap * 3)) / 2
    for i, (t, d) in enumerate(flow):
        xx = x0 + i * (bw + gap)
        if i:
            arrow_r(s, xx - gap + 0.05, 2.35, w=0.4, h=0.4, color=CYAN)
        box(s, xx, 1.9, bw, 1.35, [(t, 18, GOLD, True, 6), (d, 13.5, SUB, False)])
    tx(s, 0.8, 3.7, 11.7, 2.6,
       [("7 项校验:完整性 / 累计勾稽 / 占其所持 / 占总股本 / 风险提示勾稽 / 表格合计行 / 原文溯源覆盖", 16, WHITE, False, 10),
        ("原则:程序是会计,AI 是解说员 —— 数字对不对,程序说了算", 17, GOLD, True, 10),
        ("实测:30 份未见公告 100% 跑通,校验项通过率 93.5%", 16, GREEN, False)])

    # P8 选题一·多智能体
    s = add_slide(prs)
    page_title(s, "选题一(3/3):质疑-裁决多智能体审查")
    roles = [("提取员", "DeepSeek", "读公告、提字段", WHITE),
             ("质疑员", "Kimi(异构模型)", "只许挑毛病,质疑须引用原文", GOLD),
             ("裁决员", "DeepSeek", "回原文亲自核对:成立→修正,不成立→驳回", GREEN)]
    for i, (t, m, d, c) in enumerate(roles):
        box(s, 0.8 + i * 4.15, 1.7, 3.75, 2.0,
            [(t, 20, c, True, 6), (m, 13, SUB, False, 6), (d, 14, WHITE, False)])
    tx(s, 0.8, 4.1, 11.7, 2.4,
       [("四道防线:质疑可能编证据→裁决员不得偏信;裁决引用的原文由程序核验真实存在;", 15, WHITE, False, 6),
        ("修正只能落在数据字段白名单;修正后自动重跑勾稽与溯源。", 15, WHITE, False, 10),
        ("实战案例:惠发食品公告自身表格与文字口径矛盾(1570万-3500万+3600万≠5170万),", 15, GOLD, False, 4),
        ("三层机制合力发现“公告印错了”;对正确样本零质疑、零误报。", 15, GOLD, False)])

    # P9 选题二·数据铁律
    s = add_slide(prs)
    page_title(s, "选题二(1/2):财报数据 —— 双源采集与“代码算数”铁律")
    box(s, 0.8, 1.8, 5.7, 4.2,
        [("双源采集", 20, CYAN, True, 10),
         ("akshare 同时取同花顺摘要与新浪指标", 16, WHITE, False, 6),
         ("两源交叉印证,留痕落盘可复核", 16, WHITE, False, 6),
         ("原始 CSV 网页上直接可查", 15, SUB, False)],
        align=PP_ALIGN.LEFT)
    box(s, 6.8, 1.8, 5.7, 4.2,
        [("代码算数铁律", 20, GREEN, True, 10),
         ("所有指标由代码对落盘数据计算", 16, WHITE, False, 6),
         ("LLM 只负责解释,从不口算", 16, WHITE, False, 6),
         ("银行业口径特殊:R1/R6/R7 自动豁免,避免误报", 15, SUB, False)],
        align=PP_ALIGN.LEFT)

    # P10 选题二·九条红旗
    s = add_slide(prs)
    page_title(s, "选题二(2/2):九条红旗规则,条条有出处")
    rows = [("R1 利润与现金流背离 / R2 非经常性损益异常 / R7 利润含金量不足", "盈利质量:应计异象", "Sloan 1996;O'glove 1987"),
            ("R3 营收与利润背离 / R4 业绩变脸 / R5 持续亏损", "财务造假识别", "Beneish 1999;Dechow 2011"),
            ("R6 高杠杆", "破产风险度量", "Altman 1968;Ohlson 1980"),
            ("R8 股东高比例质押 / R9 质押到期压力", "股权质押地雷", "谢德仁等 2016《管理世界》")]
    for i, (a, b, c) in enumerate(rows):
        box(s, 0.8, 1.65 + i * 1.3, 11.7, 1.1,
            [(a, 15.5, WHITE, True, 3), (b + "  |  " + c, 13.5, GOLD, False)],
            align=PP_ALIGN.LEFT, fill=PANEL2 if i == 3 else PANEL,
            line=GOLD if i == 3 else LINE)
    tx(s, 0.8, 6.95, 11.7, 0.5,
       [("金色行由选题一公告提取结果直接驱动 —— 文献支撑 + 数据联动", 14, SUB, False)])

    # P11 实验设置
    s = add_slide(prs)
    page_title(s, "实验设置:真公告、真财报、可复现")
    tx(s, 0.8, 1.8, 11.7, 4.6,
       [("· 评测集:30 份 2026 年 6-9 月巨潮资讯网真实公告,系统从未见过", 17, WHITE, False, 12),
        ("· 预警案例:康美药业(爆雷回溯)、贵州茅台(健康对照)、贵阳银行(行业豁免)、三湘印象(质押联动)", 17, WHITE, False, 12),
        ("· 模型:DeepSeek(temperature=0,可复现)+ Kimi(异构审查),均为国产模型", 17, WHITE, False, 12),
        ("· 消融对照:同一批公告跑“纯规则路径”,量化 AI+校验组合的贡献", 17, WHITE, False, 12),
        ("· 全部数据、日志、审查卷宗随代码包开放", 17, GOLD, False)])

    # P12 实验结果
    s = add_slide(prs)
    page_title(s, "实验结果:30 份盲测")
    nums = [("100%", "跑通率\n30/30", GREEN), ("93.5%", "校验项通过率\n217/232", GOLD),
            ("91.5%", "平均溯源覆盖率\n14 份达 100%", CYAN), ("1.67 条", "平均每份提取记录\n纯规则仅 0.40 条", GREEN)]
    cw, gap = 2.85, 0.35
    gx = (W - (cw * 4 + gap * 3)) / 2
    for i, (n, d, c) in enumerate(nums):
        box(s, gx + i * (cw + gap), 2.0, cw, 2.6, [(n, 38, c, True, 10), (d, 14, WHITE, False)])
    tx(s, 0.8, 5.1, 11.7, 1.4,
       [("文档级全过:本系统 18/30(60%),纯规则 6/30(20%)—— AI+校验,缺一不可", 17, GOLD, True, 8),
        ("消融结论:LLM 提供召回主体,校验-重试与三层校验保证正确性", 16, WHITE, False)])

    # P13 康美回溯
    s = add_slide(prs)
    page_title(s, "回溯验证:康美药业提前 1-2 年报警")
    rows = [("2017 年报", "R2 非经常性损益异常", "净利润 21.50 亿 vs 扣非 40.28 亿", "早约 1 年"),
            ("2018 年报", "R1 利润与现金流背离", "净利润 +3.74 亿,每股经营现金流 -0.64 元", "早约 1 年"),
            ("2018 年报", "R4 业绩变脸", "净利润同比 -83%", "早约 1 年")]
    for i, (a, b, c, d) in enumerate(rows):
        box(s, 0.8, 1.7 + i * 1.35, 11.7, 1.15,
            [(a + "  |  " + b + "  |  " + d, 16, GOLD, True, 3), (c, 14, WHITE, False)],
            align=PP_ALIGN.LEFT)
    tx(s, 0.8, 5.9, 11.7, 0.9,
       [("对照:贵州茅台等健康公司零误报;贵阳银行自动适用行业豁免", 16, GREEN, False)])

    # P14 创新点
    s = add_slide(prs)
    page_title(s, "创新点:站在冠军方案的肩膀上,再进一步")
    tx(s, 0.8, 1.7, 11.7, 1.0,
       [("借鉴 AFAC2025 冠军项目 FinSight(人大 NLPIR)的批评循环与留痕思想,针对“提取+预警”场景做了适配性创新:", 16, SUB, False)])
    inno = [("质疑-裁决审数据", "把 FinSight 的“VLM 审图”搬到“审数据字段”;异构模型质疑、程序核验引用真实性、修正白名单"),
            ("溯源覆盖率指标", "不止标来源,而是每个字段值回原文逐页搜索,找不到显式标记“未溯源”"),
            ("代码算数铁律", "比 CAVM 更极端:完全不让 LLM 碰计算,程序套公式、AI 只解释"),
            ("双选题咬合", "选题一的提取结果直接驱动选题二 R8/R9,一条数据链贯通两个赛题")]
    for i, (t, d) in enumerate(inno):
        box(s, 0.8, 2.75 + i * 1.05, 11.7, 0.92,
            [(t, 16, GOLD, True, 2), (d, 13.5, WHITE, False)], align=PP_ALIGN.LEFT)

    # P15 总结
    s = add_slide(prs)
    page_title(s, "总结:能力 + 纪律")
    tx(s, 0.8, 1.8, 11.7, 4.4,
       [("提取得准 —— 30 份盲测校验项通过率 93.5%,溯源覆盖率 91.5%", 19, WHITE, False, 14),
        ("审得出来 —— 程序勾稽与多智能体合力发现公告自身的数据矛盾,正确样本零误报", 19, WHITE, False, 14),
        ("报得够早 —— 康美药业红色预警比爆雷早约 1-2 年,健康公司零误报", 19, WHITE, False, 14),
        ("把大模型的能力,约束在金融场景要求的严谨性之内", 21, GOLD, True)])

    # P16 结尾
    s = add_slide(prs)
    tx(s, 0, 1.4, W, 1.0, [("公告进,预警出", 44, GOLD, True)], align=PP_ALIGN.CENTER)
    tx(s, 0, 2.6, W, 0.5, [("在线演示,欢迎亲手体验", 20, WHITE, False)], align=PP_ALIGN.CENTER)
    if os.path.exists(QR_PNG):
        s.shapes.add_picture(QR_PNG, Inches(W / 2 - 0.95), Inches(3.2), Inches(1.9), Inches(1.9))
    tx(s, 0, 5.25, W, 0.45, [(DEMO_URL, 14, SUB, False)], align=PP_ALIGN.CENTER)
    tx(s, 0, 6.3, W, 0.5, [("恳请各位评委老师批评指正", 16, SUB, False)], align=PP_ALIGN.CENTER)

    path = os.path.join(OUT_DIR, "驼研信鉴_项目介绍版.pptx")
    prs.save(path)
    return path


if __name__ == "__main__":
    p1, secs = build_video()
    print("[OK] video deck ->", p1, "| slides: 10 | total:", secs, "s")
    p2 = build_intro()
    print("[OK] intro deck ->", p2, "| slides: 16")
